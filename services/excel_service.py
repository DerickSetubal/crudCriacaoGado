import io
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Cores do Tema Agro-Tech Profissional
HEADER_FILL = PatternFill(start_color="1B4D3E", end_color="1B4D3E", fill_type="solid") # Verde Floresta Escuro
SUBHEADER_FILL = PatternFill(start_color="2D6A4F", end_color="2D6A4F", fill_type="solid") # Verde Musgo
ZEBRA_FILL = PatternFill(start_color="F4F7F4", end_color="F4F7F4", fill_type="solid") # Fundo alternado suave
ACCENT_FILL = PatternFill(start_color="D8F3DC", end_color="D8F3DC", fill_type="solid") # Verde claro destaque
KPI_FILL = PatternFill(start_color="E9F5ED", end_color="E9F5ED", fill_type="solid")
TOTAL_FILL = PatternFill(start_color="D1E7DD", end_color="D1E7DD", fill_type="solid")

HEADER_FONT = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
SUBHEADER_FONT = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Segoe UI", size=16, bold=True, color="1B4D3E")
SUBTITLE_FONT = Font(name="Segoe UI", size=10, italic=True, color="4B5563")
BOLD_FONT = Font(name="Segoe UI", size=10, bold=True, color="1F2937")
REGULAR_FONT = Font(name="Segoe UI", size=10, color="1F2937")
KPI_TITLE_FONT = Font(name="Segoe UI", size=9, bold=True, color="2D6A4F")
KPI_VALUE_FONT = Font(name="Segoe UI", size=14, bold=True, color="1B4D3E")

THIN_BORDER_SIDE = Side(border_style="thin", color="D1D5DB")
THIN_BORDER = Border(left=THIN_BORDER_SIDE, right=THIN_BORDER_SIDE, top=THIN_BORDER_SIDE, bottom=THIN_BORDER_SIDE)
TOTAL_TOP_DOUBLE = Border(top=Side(border_style="thin", color="1B4D3E"), bottom=Side(border_style="double", color="1B4D3E"))

ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")


def ajustar_largura_colunas(ws):
    """Ajusta automaticamente a largura das colunas com margem de respiro."""
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            # Ignora células mescladas no cálculo do comprimento para não esticar colunas
            if cell.coordinate in ws.merged_cells:
                continue
            if cell.value is not None:
                valor_str = str(cell.value)
                if len(valor_str) > max_len:
                    max_len = len(valor_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)


def gerar_planilha_consolidada(lotes, manejos=None, pastos=None, titulo_relatorio="Relatório Consolidado de Pecuária"):
    """
    Gera um workbook Excel (.xlsx) altamente profissional contendo:
    - Aba 1: Resumo Executivo & KPIs
    - Aba 2: Lotes de Gado (Entrada & Evolução)
    - Aba 3: Histórico Trimestral de Manejos (Pesagem, Vacinas, Remédios)
    - Aba 4: Distribuição por Pasto/Cercado
    """
    wb = openpyxl.Workbook()
    # Remove aba padrão inicial
    ws_default = wb.active
    
    # =========================================================================
    # ABA 1: RESUMO EXECUTIVO & KPIs
    # =========================================================================
    ws_kpi = wb.create_sheet(title="Resumo Executivo")
    ws_kpi.views.sheetView[0].showGridLines = True
    
    # Cabeçalho da Empresa / Fazenda
    ws_kpi.merge_cells("A1:G1")
    ws_kpi["A1"] = "SISTEMA DE GESTÃO PECUÁRIA - FAZENDA MODELO"
    ws_kpi["A1"].font = TITLE_FONT
    ws_kpi["A1"].alignment = ALIGN_LEFT
    
    ws_kpi.merge_cells("A2:G2")
    ws_kpi["A2"] = f"{titulo_relatorio} | Gerado em: {datetime.now().strftime('%d/%m/%Y às %H:%M')}"
    ws_kpi["A2"].font = SUBTITLE_FONT
    ws_kpi["A2"].alignment = ALIGN_LEFT
    
    # Métricas Consolidadas
    total_lotes_count = len(lotes)
    total_cabecas_count = sum(l.quantidade_cabecas for l in lotes) if lotes else 0
    investimento_total = sum(l.valor_total_aquisicao for l in lotes) if lotes else 0
    
    pesos_atuais = [l.peso_atual_medio for l in lotes if l.peso_atual_medio]
    peso_medio_geral = (sum(pesos_atuais) / len(pesos_atuais)) if pesos_atuais else 0.0
    
    ganhos_acumulados = [l.ganho_peso_total_acumulado for l in lotes]
    ganho_medio_geral = (sum(ganhos_acumulados) / len(ganhos_acumulados)) if ganhos_acumulados else 0.0
    
    arrobas_totais_rebanho = sum(l.arrobas_atuais_totais for l in lotes) if lotes else 0.0

    # Cards de KPIs no Excel
    kpis = [
        ("TOTAL DE ANIMAIS", f"{total_cabecas_count:,} cab".replace(",", "."), "B4:C5"),
        ("LOTES ATIVOS", f"{total_lotes_count} lotes", "D4:D5"),
        ("PESO MÉDIO ATUAL", f"{peso_medio_geral:.1f} kg", "E4:E5"),
        ("GANHO MÉDIO TOTAL", f"+{ganho_medio_geral:.1f} kg", "F4:F5"),
        ("TOTAL DE ARROBAS (@)", f"{arrobas_totais_rebanho:,.1f} @".replace(",", "X").replace(".", ",").replace("X", "."), "G4:G5"),
    ]
    
    # Tabela Rápida de Indicadores na Aba Resumo
    ws_kpi["A4"] = "INDICADOR"
    ws_kpi["A4"].font = SUBHEADER_FONT
    ws_kpi["A4"].fill = SUBHEADER_FILL
    ws_kpi["A4"].alignment = ALIGN_LEFT
    
    ws_kpi["B4"] = "VALOR CONSOLIDADO"
    ws_kpi["B4"].font = SUBHEADER_FONT
    ws_kpi["B4"].fill = SUBHEADER_FILL
    ws_kpi["B4"].alignment = ALIGN_RIGHT
    
    linhas_kpi = [
        ("Total de Animais Ativos (Cabeças)", total_cabecas_count, '#,##0 "cab"'),
        ("Total de Lotes Cadastrados", total_lotes_count, '#,##0 "lotes"'),
        ("Peso Médio Atual do Rebanho (kg/cab)", peso_medio_geral, '#,##0.0 "kg"'),
        ("Ganho Médio Acumulado por Animal (kg)", ganho_medio_geral, '+#,##0.0 "kg"'),
        ("Total de Arrobas Produzidas (@)", arrobas_totais_rebanho, '#,##0.0 "@"'),
        ("Investimento Total em Aquisição (R$)", investimento_total, 'R$ #,##0.00'),
    ]
    
    curr_r = 5
    for rotulo, val, fmt in linhas_kpi:
        ws_kpi[f"A{curr_r}"] = rotulo
        ws_kpi[f"A{curr_r}"].font = BOLD_FONT
        ws_kpi[f"A{curr_r}"].border = THIN_BORDER
        if curr_r % 2 == 0:
            ws_kpi[f"A{curr_r}"].fill = ZEBRA_FILL
            
        c = ws_kpi[f"B{curr_r}"]
        c.value = val
        c.font = BOLD_FONT
        c.number_format = fmt
        c.alignment = ALIGN_RIGHT
        c.border = THIN_BORDER
        if curr_r % 2 == 0:
            c.fill = ZEBRA_FILL
        curr_r += 1

    # Regras e Notas explicativas no Resumo
    curr_r += 2
    ws_kpi[f"A{curr_r}"] = "NOTAS OPERACIONAIS E DIRETRIZES DO MANEJO:"
    ws_kpi[f"A{curr_r}"].font = BOLD_FONT
    curr_r += 1
    notas = [
        "1. Ciclo de Manejo Trimestral: Repesagem, vacinação e medicação obrigatórias a cada 90 dias.",
        "2. Arroba Viva: Cálculo padrão adotado de 30 kg vivo = 1 @ (equivalente a 15 kg de carcaça / rendimento 50%).",
        "3. Monitoramento de GMD: Ganho Médio Diário aferido entre a pesagem de entrada e os ciclos subsequentes.",
        "4. Rotação de Pastagem: Mantenha a alocação atualizada a cada manejo para evitar degradação de cercados."
    ]
    for n in notas:
        ws_kpi[f"A{curr_r}"] = n
        ws_kpi[f"A{curr_r}"].font = REGULAR_FONT
        curr_r += 1

    ajustar_largura_colunas(ws_kpi)

    # =========================================================================
    # ABA 2: LOTES DE GADO (ENTRADA & SITUAÇÃO ATUAL)
    # =========================================================================
    ws_lotes = wb.create_sheet(title="Lotes de Gado")
    ws_lotes.views.sheetView[0].showGridLines = True
    
    ws_lotes.merge_cells("A1:M1")
    ws_lotes["A1"] = "CONTROLE CONSOLIDADO DE LOTES DE GADO"
    ws_lotes["A1"].font = TITLE_FONT
    ws_lotes["A1"].alignment = ALIGN_LEFT
    
    headers_lotes = [
        "Código Lote", "Lote Origem", "Data Entrada", "Cabeças",
        "Pasto/Cercado", "Raça", "Peso Entr. (kg)", "Valor @ Pago (R$)",
        "Invest. Total (R$)", "Peso Atual (kg)", "Ganho Acum. (kg)", "Arrobas Atuais (@)", "Status Manejo"
    ]
    
    row_num = 3
    for col_idx, header in enumerate(headers_lotes, start=1):
        cell = ws_lotes.cell(row=row_num, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = ALIGN_CENTER
        cell.border = THIN_BORDER
        
    start_row_lotes = 4
    for idx, lote in enumerate(lotes, start=start_row_lotes):
        fill_atual = ZEBRA_FILL if idx % 2 == 0 else PatternFill(fill_type=None)
        
        ws_lotes.cell(row=idx, column=1, value=lote.codigo).alignment = ALIGN_CENTER
        ws_lotes.cell(row=idx, column=2, value=lote.lote_origem).alignment = ALIGN_LEFT
        
        c_dt = ws_lotes.cell(row=idx, column=3, value=lote.data_chegada.strftime("%d/%m/%Y") if lote.data_chegada else "")
        c_dt.alignment = ALIGN_CENTER
        
        c_cab = ws_lotes.cell(row=idx, column=4, value=lote.quantidade_cabecas)
        c_cab.alignment = ALIGN_RIGHT
        c_cab.number_format = '#,##0'
        
        ws_lotes.cell(row=idx, column=5, value=lote.pasto_atual).alignment = ALIGN_LEFT
        ws_lotes.cell(row=idx, column=6, value=lote.raca).alignment = ALIGN_LEFT
        
        c_pe = ws_lotes.cell(row=idx, column=7, value=lote.peso_chegada_medio)
        c_pe.alignment = ALIGN_RIGHT
        c_pe.number_format = '#,##0.0'
        
        c_va = ws_lotes.cell(row=idx, column=8, value=lote.valor_arroba_pago)
        c_va.alignment = ALIGN_RIGHT
        c_va.number_format = 'R$ #,##0.00'
        
        c_tot = ws_lotes.cell(row=idx, column=9, value=lote.valor_total_aquisicao)
        c_tot.alignment = ALIGN_RIGHT
        c_tot.number_format = 'R$ #,##0.00'
        
        c_pa = ws_lotes.cell(row=idx, column=10, value=lote.peso_atual_medio)
        c_pa.alignment = ALIGN_RIGHT
        c_pa.number_format = '#,##0.0'
        
        c_ganho = ws_lotes.cell(row=idx, column=11, value=lote.ganho_peso_total_acumulado)
        c_ganho.alignment = ALIGN_RIGHT
        c_ganho.number_format = '+#,##0.0'
        
        c_arr = ws_lotes.cell(row=idx, column=12, value=lote.arrobas_atuais_totais)
        c_arr.alignment = ALIGN_RIGHT
        c_arr.number_format = '#,##0.0'
        
        status_info = lote.status_manejo
        c_st = ws_lotes.cell(row=idx, column=13, value=status_info.get('status', 'Em dia'))
        c_st.alignment = ALIGN_CENTER
        if status_info.get('status') == 'Atrasado':
            c_st.font = Font(name="Segoe UI", size=10, bold=True, color="991B1B")
        elif status_info.get('status') == 'Próximo':
            c_st.font = Font(name="Segoe UI", size=10, bold=True, color="92400E")
        else:
            c_st.font = Font(name="Segoe UI", size=10, color="065F46")
        
        for c_idx in range(1, 14):
            cel = ws_lotes.cell(row=idx, column=c_idx)
            cel.border = THIN_BORDER
            if fill_atual.fill_type:
                cel.fill = fill_atual
                
    end_row_lotes = len(lotes) + start_row_lotes - 1
    if end_row_lotes >= start_row_lotes:
        tot_row = end_row_lotes + 1
        ws_lotes.cell(row=tot_row, column=1, value="TOTALIZADORES / MÉDIAS").alignment = ALIGN_LEFT
        ws_lotes.cell(row=tot_row, column=1).font = BOLD_FONT
        
        # Fórmula Excel de Soma de Cabeças
        c_tot_cab = ws_lotes.cell(row=tot_row, column=4, value=f"=SUM(D{start_row_lotes}:D{end_row_lotes})")
        c_tot_cab.alignment = ALIGN_RIGHT
        c_tot_cab.font = BOLD_FONT
        c_tot_cab.number_format = '#,##0'
        
        # Média de Peso Entrada
        c_med_pe = ws_lotes.cell(row=tot_row, column=7, value=f"=AVERAGE(G{start_row_lotes}:G{end_row_lotes})")
        c_med_pe.alignment = ALIGN_RIGHT
        c_med_pe.font = BOLD_FONT
        c_med_pe.number_format = '#,##0.0'
        
        # Média Valor @
        c_med_va = ws_lotes.cell(row=tot_row, column=8, value=f"=AVERAGE(H{start_row_lotes}:H{end_row_lotes})")
        c_med_va.alignment = ALIGN_RIGHT
        c_med_va.font = BOLD_FONT
        c_med_va.number_format = 'R$ #,##0.00'
        
        # Soma Investimento Total
        c_soma_inv = ws_lotes.cell(row=tot_row, column=9, value=f"=SUM(I{start_row_lotes}:I{end_row_lotes})")
        c_soma_inv.alignment = ALIGN_RIGHT
        c_soma_inv.font = BOLD_FONT
        c_soma_inv.number_format = 'R$ #,##0.00'
        
        # Média Peso Atual
        c_med_pa = ws_lotes.cell(row=tot_row, column=10, value=f"=AVERAGE(J{start_row_lotes}:J{end_row_lotes})")
        c_med_pa.alignment = ALIGN_RIGHT
        c_med_pa.font = BOLD_FONT
        c_med_pa.number_format = '#,##0.0'
        
        # Média Ganho Acumulado
        c_med_ga = ws_lotes.cell(row=tot_row, column=11, value=f"=AVERAGE(K{start_row_lotes}:K{end_row_lotes})")
        c_med_ga.alignment = ALIGN_RIGHT
        c_med_ga.font = BOLD_FONT
        c_med_ga.number_format = '+#,##0.0'
        
        # Soma Arrobas Atuais
        c_soma_arr = ws_lotes.cell(row=tot_row, column=12, value=f"=SUM(L{start_row_lotes}:L{end_row_lotes})")
        c_soma_arr.alignment = ALIGN_RIGHT
        c_soma_arr.font = BOLD_FONT
        c_soma_arr.number_format = '#,##0.0'
        
        for c_idx in range(1, 14):
            cel = ws_lotes.cell(row=tot_row, column=c_idx)
            cel.fill = TOTAL_FILL
            cel.border = TOTAL_TOP_DOUBLE

    ajustar_largura_colunas(ws_lotes)

    # =========================================================================
    # ABA 3: HISTÓRICO DE MANEJOS TRIMESTRAIS (Pesagem, Vacinas, Remédios)
    # =========================================================================
    ws_man = wb.create_sheet(title="Histórico Manejos (3 Meses)")
    ws_man.views.sheetView[0].showGridLines = True
    
    ws_man.merge_cells("A1:K1")
    ws_man["A1"] = "CICLO DE MANEJO TRIMESTRAL (PESAGEM, VACINAÇÃO E REMÉDIOS)"
    ws_man["A1"].font = TITLE_FONT
    ws_man["A1"].alignment = ALIGN_LEFT
    
    headers_man = [
        "Lote", "Ciclo", "Data Manejo", "Peso Anterior (kg)",
        "Peso Aferido (kg)", "Ganho Período (kg)", "GMD (kg/dia)",
        "Vacinas Aplicadas", "Remédios / Vermífugos", "Pasto Destino", "Observações"
    ]
    
    row_num = 3
    for col_idx, header in enumerate(headers_man, start=1):
        cell = ws_man.cell(row=row_num, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = ALIGN_CENTER
        cell.border = THIN_BORDER

    # Obter todos os manejos ordenados
    lista_manejos = []
    if manejos is not None:
        lista_manejos = manejos
    else:
        for lote in lotes:
            lista_manejos.extend(lote.manejos)
    
    lista_manejos.sort(key=lambda m: (m.data_manejo, m.lote_id), reverse=True)
    
    start_row_man = 4
    for idx, man in enumerate(lista_manejos, start=start_row_man):
        fill_atual = ZEBRA_FILL if idx % 2 == 0 else PatternFill(fill_type=None)
        
        ws_man.cell(row=idx, column=1, value=man.lote.codigo if man.lote else f"Lote #{man.lote_id}").alignment = ALIGN_CENTER
        ws_man.cell(row=idx, column=2, value=f"{man.ciclo_numero}º Manejo").alignment = ALIGN_CENTER
        
        c_dt = ws_man.cell(row=idx, column=3, value=man.data_manejo.strftime("%d/%m/%Y") if man.data_manejo else "")
        c_dt.alignment = ALIGN_CENTER
        
        c_pant = ws_man.cell(row=idx, column=4, value=man.peso_anterior_kg)
        c_pant.alignment = ALIGN_RIGHT
        c_pant.number_format = '#,##0.0'
        
        c_pat = ws_man.cell(row=idx, column=5, value=man.peso_medio_kg)
        c_pat.alignment = ALIGN_RIGHT
        c_pat.number_format = '#,##0.0'
        
        c_ganho = ws_man.cell(row=idx, column=6, value=man.ganho_peso_kg)
        c_ganho.alignment = ALIGN_RIGHT
        c_ganho.number_format = '+#,##0.0'
        
        c_gmd = ws_man.cell(row=idx, column=7, value=man.gmd_kg)
        c_gmd.alignment = ALIGN_RIGHT
        c_gmd.number_format = '0.000'
        
        ws_man.cell(row=idx, column=8, value=man.vacinas or "Nenhuma").alignment = ALIGN_LEFT
        ws_man.cell(row=idx, column=9, value=man.remedios or "Nenhum").alignment = ALIGN_LEFT
        ws_man.cell(row=idx, column=10, value=man.pasto_destino or "").alignment = ALIGN_LEFT
        ws_man.cell(row=idx, column=11, value=man.observacoes or "").alignment = ALIGN_LEFT
        
        for c_idx in range(1, 12):
            cel = ws_man.cell(row=idx, column=c_idx)
            cel.border = THIN_BORDER
            if fill_atual.fill_type:
                cel.fill = fill_atual
                
    end_row_man = len(lista_manejos) + start_row_man - 1
    if end_row_man >= start_row_man:
        tot_row_m = end_row_man + 1
        ws_man.cell(row=tot_row_m, column=1, value="MÉDIAS DO HISTÓRICO").alignment = ALIGN_LEFT
        ws_man.cell(row=tot_row_m, column=1).font = BOLD_FONT
        
        c_med_paf = ws_man.cell(row=tot_row_m, column=5, value=f"=AVERAGE(E{start_row_man}:E{end_row_man})")
        c_med_paf.alignment = ALIGN_RIGHT
        c_med_paf.font = BOLD_FONT
        c_med_paf.number_format = '#,##0.0'
        
        c_med_g = ws_man.cell(row=tot_row_m, column=6, value=f"=AVERAGE(F{start_row_man}:F{end_row_man})")
        c_med_g.alignment = ALIGN_RIGHT
        c_med_g.font = BOLD_FONT
        c_med_g.number_format = '+#,##0.0'
        
        c_med_gmd = ws_man.cell(row=tot_row_m, column=7, value=f"=AVERAGE(G{start_row_man}:G{end_row_man})")
        c_med_gmd.alignment = ALIGN_RIGHT
        c_med_gmd.font = BOLD_FONT
        c_med_gmd.number_format = '0.000'
        
        for c_idx in range(1, 12):
            cel = ws_man.cell(row=tot_row_m, column=c_idx)
            cel.fill = TOTAL_FILL
            cel.border = TOTAL_TOP_DOUBLE

    ajustar_largura_colunas(ws_man)

    # =========================================================================
    # ABA 4: DISTRIBUIÇÃO POR PASTO / CERCADO
    # =========================================================================
    ws_pastos = wb.create_sheet(title="Pastos & Cercados")
    ws_pastos.views.sheetView[0].showGridLines = True
    
    ws_pastos.merge_cells("A1:G1")
    ws_pastos["A1"] = "DISTRIBUIÇÃO DE LOTES E ANIMAIS POR PASTO/CERCADO"
    ws_pastos["A1"].font = TITLE_FONT
    ws_pastos["A1"].alignment = ALIGN_LEFT
    
    headers_pastos = [
        "Pasto/Cercado", "Tipo da Pastagem", "Área (ha)", "Capacidade (cab)",
        "Ocupação Atual (cab)", "Lotes Presentes", "Status"
    ]
    
    row_num = 3
    for col_idx, header in enumerate(headers_pastos, start=1):
        cell = ws_pastos.cell(row=row_num, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = ALIGN_CENTER
        cell.border = THIN_BORDER
        
    # Agrupar lotes por pasto
    pasto_map = {}
    if pastos:
        for p in pastos:
            pasto_map[p.nome] = {
                'obj': p,
                'cabecas': 0,
                'lotes': []
            }
            
    for l in lotes:
        nome_p = l.pasto_atual or "Não Definido"
        if nome_p not in pasto_map:
            pasto_map[nome_p] = {
                'obj': None,
                'cabecas': 0,
                'lotes': []
            }
        pasto_map[nome_p]['cabecas'] += l.quantidade_cabecas
        pasto_map[nome_p]['lotes'].append(l.codigo)
        
    start_row_p = 4
    for idx, (p_nome, p_info) in enumerate(pasto_map.items(), start=start_row_p):
        fill_atual = ZEBRA_FILL if idx % 2 == 0 else PatternFill(fill_type=None)
        p_obj = p_info['obj']
        
        ws_pastos.cell(row=idx, column=1, value=p_nome).alignment = ALIGN_LEFT
        ws_pastos.cell(row=idx, column=2, value=p_obj.tipo_pasto if p_obj else "Brachiaria").alignment = ALIGN_LEFT
        
        c_ar = ws_pastos.cell(row=idx, column=3, value=p_obj.area_hectares if p_obj else 10.0)
        c_ar.alignment = ALIGN_RIGHT
        c_ar.number_format = '#,##0.0 "ha"'
        
        c_cap = ws_pastos.cell(row=idx, column=4, value=p_obj.capacidade_cabecas if p_obj else 50)
        c_cap.alignment = ALIGN_RIGHT
        c_cap.number_format = '#,##0'
        
        c_ocup = ws_pastos.cell(row=idx, column=5, value=p_info['cabecas'])
        c_ocup.alignment = ALIGN_RIGHT
        c_ocup.number_format = '#,##0'
        
        ws_pastos.cell(row=idx, column=6, value=", ".join(p_info['lotes']) or "Nenhum").alignment = ALIGN_LEFT
        
        capacidade = p_obj.capacidade_cabecas if p_obj else 50
        status_ocupacao = "Disponível"
        if p_info['cabecas'] >= capacidade:
            status_ocupacao = "Capacidade Máxima"
        elif p_info['cabecas'] > 0:
            status_ocupacao = "Em Uso"
        else:
            status_ocupacao = "Vazio / Descanso"
            
        c_st_p = ws_pastos.cell(row=idx, column=7, value=status_ocupacao)
        c_st_p.alignment = ALIGN_CENTER
        
        for c_idx in range(1, 8):
            cel = ws_pastos.cell(row=idx, column=c_idx)
            cel.border = THIN_BORDER
            if fill_atual.fill_type:
                cel.fill = fill_atual

    ajustar_largura_colunas(ws_pastos)

    # Remove a aba default se ainda existir
    if ws_default in wb.worksheets and len(wb.worksheets) > 1:
        wb.remove(ws_default)
        
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

